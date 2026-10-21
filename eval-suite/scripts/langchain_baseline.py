# -*- coding: utf-8 -*-
"""Baseline A (2026-10-17): the standard LangChain-recipe RAG pipeline.

What a competent engineer builds in a day with the default framework tools —
the industry "reference" against which our evidence layer is compared:

  load (generic parsers) → RecursiveCharacterTextSplitter(800/120)
  → bge-m3 embeddings (local) → BM25(jieba/en tokens) + vector cosine
  → 50/50 RRF ensemble → top-8 chunks → generic RAG prompt → same gateway LLM.

Same source documents (111 SP files), same benchmark (v2_289), same judge.
Retrieved chunks enter the snapshot as evidence (doc_id = original doc_uid via
the export map) so the judge's groundedness dims stay comparable.

Usage:
  ../llm-corpus-builder/.venv/Scripts/python.exe scripts/langchain_baseline.py \
      --eval <bench> --out <snapshot.json> [--workers 10]
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import httpx
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from evalsuite import builder_root  # noqa: E402

DOCS_DIR = ROOT / "data" / "external_baselines" / "sp_source_docs"
MAP = ROOT / "data" / "external_baselines" / "sp_source_docs_map.json"

PROMPT = """你是高校教务问答助手。仅根据下面的参考资料回答问题。
规则：
1. 资料中没有答案时，明确说明"参考资料中没有该信息"，禁止编造数字、日期、课程、政策
2. 回答末尾用 [1][2] 标注引用了哪些资料
3. 中文问题用中文回答，英文问题用英文回答

参考资料：
{context}

问题：{question}"""


def _html_text(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text(" ", strip=True)


def _pdf_text(path: Path) -> str:
    import pymupdf

    parts = []
    with pymupdf.open(path) as doc:
        for page in doc:
            parts.append(page.get_text(sort=True))
    return "\n".join(parts)


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    try:
        for sheet in wb.worksheets:
            parts.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    parts.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(parts)


LOADERS = {".html": _html_text, ".pdf": _pdf_text, ".docx": _docx_text, ".xlsx": _xlsx_text}


def build_index():
    """Load → chunk → embed → BM25 corpus. Runs once."""
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from rank_bm25 import BM25Okapi
    import jieba
    from sentence_transformers import SentenceTransformer
    import numpy as np

    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=120)
    doc_map = json.loads(MAP.read_text(encoding="utf-8"))
    chunks: list[dict] = []
    for path in sorted(DOCS_DIR.iterdir()):
        loader = LOADERS.get(path.suffix.lower())
        if not loader:
            continue
        text = loader(path)
        if not text.strip():
            continue
        for piece in splitter.split_text(text):
            chunks.append({"file": path.name, "doc_uid": doc_map.get(path.name, ""), "text": piece})
    print(f"[lc] {len(chunks)} chunks over {len(set(c['file'] for c in chunks))} docs", flush=True)

    def tokens(text: str) -> list[str]:
        en = re.findall(r"[A-Za-z0-9]+", text)
        return list(jieba.cut_for_search(text)) + en

    bm25 = BM25Okapi([tokens(c["text"]) for c in chunks])
    model = SentenceTransformer(str(builder_root() / "models" / "embed" / "bge-m3"))
    matrix = model.encode([c["text"] for c in chunks], batch_size=32, normalize_embeddings=True, show_progress_bar=False)
    matrix = np.asarray(matrix)
    print("[lc] embeddings ready", flush=True)
    return chunks, bm25, tokens, model, matrix


def retrieve(question: str, k: int, bundle) -> list[dict]:
    chunks, bm25, tokens, model, matrix = bundle
    import numpy as np

    q_tokens = tokens(question)
    bm25_scores = bm25.get_scores(q_tokens)
    bm25_rank = sorted(range(len(chunks)), key=lambda i: -bm25_scores[i])[:200]
    qvec = model.encode([question], normalize_embeddings=True)[0]
    cos = matrix @ qvec
    dense_rank = sorted(range(len(chunks)), key=lambda i: -cos[i])[:200]
    fused: dict[int, float] = {}
    for rank, i in enumerate(bm25_rank, 1):
        fused[i] = fused.get(i, 0.0) + 1.0 / (60 + rank)
    for rank, i in enumerate(dense_rank, 1):
        fused[i] = fused.get(i, 0.0) + 1.0 / (60 + rank)
    top = sorted(fused, key=lambda i: -fused[i])[:k]
    return [{**chunks[i], "score": float(fused[i])} for i in top]


def ask(prompt: str) -> str:
    base = os.environ.get("LLM_BASE_URL", "https://your-gateway.example.com/v1")
    key = os.environ.get("LLM_API_KEY") or os.environ.get("QUESTIONGEN_API_KEY", "")
    model = os.environ.get("LLM_MODEL", "deepseek-v4-flash-780k-local")
    body = {"model": model, "messages": [{"role": "user", "content": prompt}],
            "temperature": 0, "max_tokens": 800, "stream": True, "thinking": {"type": "disabled"}}
    for attempt in range(4):
        try:
            text = ""
            with httpx.stream("POST", f"{base}/chat/completions", headers={"Authorization": f"Bearer {key}"},
                              json=body, timeout=300) as resp:
                if resp.status_code == 429:
                    time.sleep(20 + attempt * 20)
                    continue
                resp.raise_for_status()
                for line in resp.iter_lines():
                    if not line.startswith("data: "):
                        continue
                    payload = line[6:]
                    if payload.strip() == "[DONE]":
                        break
                    try:
                        text += json.loads(payload)["choices"][0]["delta"].get("content") or ""
                    except Exception:
                        continue
            return text.strip()
        except Exception:
            if attempt == 3:
                return ""
            time.sleep(5 + attempt * 10)
    return ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--eval", type=Path, default=ROOT / "data" / "benchmarks" / "corpus-eval" / "benchmark_superset_v2_289.jsonl")
    parser.add_argument("--out", type=Path, default=ROOT / "data" / "benchmarks" / "corpus-eval" / "results_e2e_product" / "ext_langchain.json")
    parser.add_argument("--today", default="2026-10-15")
    parser.add_argument("--workers", type=int, default=10)
    parser.add_argument("--k", type=int, default=8)
    args = parser.parse_args()

    bundle = build_index()
    bench = [json.loads(l) for l in open(args.eval, encoding="utf-8")]
    done: dict[str, dict] = {}
    if args.out.is_file():
        done = {r["id"]: r for r in json.loads(args.out.read_text(encoding="utf-8"))["rows"]}
        print(f"resume: {len(done)}")
    todo = [b for b in bench if b["id"] not in done]
    rows: list[dict] = []

    def run(item):
        started = time.perf_counter()
        hits = retrieve(item["question"], args.k, bundle)
        context = "\n\n".join(f"[{i}] {h['text'][:1200]}" for i, h in enumerate(hits, 1))
        answer = ask(PROMPT.format(context=context, question=item["question"]))
        return item, hits, answer, (time.perf_counter() - started) * 1000

    def write():
        args.out.parent.mkdir(parents=True, exist_ok=True)
        snapshot = {"meta": {"label": "ext-langchain", "url": "in-process + gateway", "n": len(rows) + len(done),
                             "providers": ["langchain-recipe"], "snapshot_version": 3,
                             "recipe": "RecursiveCharacterTextSplitter 800/120 + bge-m3 + BM25(jieba)/cosine RRF + top8"},
                    "rows": rows + list(done.values())}
        args.out.write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")

    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = [pool.submit(run, item) for item in todo]
        for n, fut in enumerate(as_completed(futures), 1):
            item, hits, answer, elapsed = fut.result()
            abstained = any(m in answer for m in ("没有该信息", "参考资料中没", "不知道", "无法确定", "not in the", "no information", "I don't know"))
            cited = sorted({int(m) for m in re.findall(r"\[(\d)\]", answer)})
            rows.append({
                "id": item["id"], "category": item.get("category", ""), "level": item.get("level"),
                "language": item.get("language"), "qshape": item.get("qshape", ""),
                "answerable": bool(item.get("gold")),
                "answer_text": answer, "answer_disposition": "evidence_refusal" if abstained else "answer",
                "abstained": abstained, "cites_evidence": bool(cited), "grounded": bool(cited),
                "numeric_fidelity": True, "latency_ms": round(elapsed, 1),
                "claims": [], "citations": [{"doc_uid": hits[i - 1]["doc_uid"], "title": hits[i - 1]["file"]} for i in cited if i <= len(hits)],
                "citation_map": {},
                "decision": {"status": "sufficient", "abstain": abstained},
                "generation": {"mode": "langchain"},
                "evidence_pack": {"evidence": [{"evidence_id": f"E{i}", "doc_id": h["doc_uid"], "chunk_id": "", "title": h["file"], "text": h["text"][:1200]} for i, h in enumerate(hits, 1)],
                                  "usable_evidence_ids": [f"E{i}" for i in range(1, len(hits) + 1)],
                                  "requirements": [], "evidence_status": "sufficient", "abstain": False,
                                  "abstain_reason": "", "coverage": 0.0, "conflicts": [], "fact_conflicts": []},
                "retrieval_trace": {"strategy": "langchain_ensemble", "queries": [], "rounds": []},
            })
            if n % 20 == 0:
                print(f"[lc] {n}/{len(todo)}", flush=True)
                write()
    write()
    print(f"done: {len(rows) + len(done)} -> {args.out}")


if __name__ == "__main__":
    main()
