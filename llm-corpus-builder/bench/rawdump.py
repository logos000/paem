"""Naive plain-text extraction for baseline corpora (P7.1).

标准 RAG 流程怎么拿文本，这里就怎么拿：PDF 逐页 ``get_text`` 拼接、DOCX
段落直读、XLSX 逐格拼接、TXT 原样。刻意**不做**结构恢复——没有表格行表头
自包含、没有标题层级、没有阅读顺序修正。基线语料不得偷用本项目的结构化
抽取（否则 Track A 测的是我们的方法，不是常规流程）。
"""

from __future__ import annotations

from pathlib import Path

__all__ = ["plain_text", "html_plain"]


def html_plain(raw: bytes) -> str:
    """Naive HTML→text the way a standard RAG pipeline does it.

    ``BeautifulSoup.get_text`` over the whole body — navigation shells, footer
    links and boilerplate included.  No reading-order fixing, no table
    self-containment, no boilerplate separation: those are exactly the corpus
    contributions Track A is measuring against.
    """
    try:
        from bs4 import BeautifulSoup

        for encoding in ("utf-8", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("utf-8", "replace")
        soup = BeautifulSoup(text, "lxml")
        for tag in soup(["script", "style", "noscript", "template"]):
            tag.decompose()
        return soup.get_text("\n")
    except Exception:
        return ""


def plain_text(path: Path, max_pages: int = 400) -> str:
    suffix = Path(path).suffix.lower()
    try:
        if suffix == ".pdf":
            import pymupdf

            parts = []
            with pymupdf.open(path) as doc:
                for number, page in enumerate(doc, 1):
                    if number > max_pages:
                        break
                    parts.append(page.get_text())
            return "\n".join(parts)
        if suffix in (".docx",):
            from docx import Document

            doc = Document(str(path))
            return "\n".join(paragraph.text for paragraph in doc.paragraphs)
        if suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook

            workbook = load_workbook(str(path), read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(cell) for cell in row if cell is not None and str(cell).strip()]
                    if cells:
                        parts.append("\t".join(cells))
            workbook.close()
            return "\n".join(parts)
        if suffix in (".txt", ".md", ".html", ".htm"):
            raw = Path(path).read_bytes()
            for encoding in ("utf-8", "gb18030"):
                try:
                    return raw.decode(encoding)
                except UnicodeDecodeError:
                    continue
            return raw.decode("utf-8", "replace")
    except Exception:
        return ""
    return ""
