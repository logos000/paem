"""Structure-preserving extraction of HTML, PDF, DOCX and XLSX into evidence blocks.

Every extractor yields :class:`RawBlock` items in reading order.  Tables are
emitted row by row with their header so a row is self-contained evidence
("Class: Seniors | Start time: 8:00 pm Dec 19th 2025") and still carries its
``table_id``/``row_index``/``cells`` for exact citation.
"""

from __future__ import annotations

from dataclasses import dataclass, field
import re
from pathlib import Path
from typing import Any, Iterable, Iterator

from bs4 import BeautifulSoup, NavigableString, Tag

from evidence.store import normalize_text

BLOCK_TAGS = {
    "p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "table", "ul", "ol", "dl", "dt", "dd",
    "pre", "blockquote", "div", "section", "article", "header", "footer", "nav", "aside",
    "form", "figure", "figcaption", "main", "center", "tr", "td", "th", "tbody", "thead", "hr", "br",
}
SKIP_TAGS = {"script", "style", "noscript", "iframe", "button", "select", "option", "svg", "canvas", "video", "audio", "object", "embed", "input", "textarea", "template"}
HEADING_PATTERNS = (
    (1, re.compile(r"^第[一二三四五六七八九十百零〇\d]+[章编]\s*\S{0,30}$")),
    (2, re.compile(r"^第[一二三四五六七八九十百零〇\d]+节\s*\S{0,30}$")),
    (2, re.compile(r"^[一二三四五六七八九十]+、\s*\S.{0,28}$")),
    (3, re.compile(r"^[（(][一二三四五六七八九十]+[）)]\s*\S.{0,28}$")),
    (2, re.compile(r"^(?:Part|Section|Chapter)\s+[IVX\d]+\b.{0,50}$", re.I)),
    (2, re.compile(r"^[A-H][.、．]\s*[A-Z一-鿿][^.。:：]{0,38}$")),
)


@dataclass
class RawBlock:
    kind: str
    text: str
    level: int = 0
    page: int | None = None
    table: dict[str, Any] | None = None
    extra: dict[str, Any] = field(default_factory=dict)


@dataclass
class Extraction:
    blocks: list[RawBlock]
    warnings: list[str] = field(default_factory=list)
    pages: int = 0
    needs_ocr_pages: list[int] = field(default_factory=list)
    links: list[dict[str, str]] = field(default_factory=list)
    # Per-page OCR bookkeeping so a missing dependency or a single bad page
    # degrades that page instead of failing the whole document.
    ocr_attempted_pages: list[int] = field(default_factory=list)
    ocr_succeeded_pages: list[int] = field(default_factory=list)
    ocr_failed_pages: list[int] = field(default_factory=list)
    ocr_dependency_missing_pages: list[int] = field(default_factory=list)


def heading_level(text: str) -> int:
    """Return a heading level for Chinese/English numbered headings, else 0."""
    text = normalize_text(text)
    if not text or len(text) > 60 or text.endswith(("。", "；", ";")):
        return 0
    for level, pattern in HEADING_PATTERNS:
        if pattern.match(text):
            return level
    return 0


# ---------------------------------------------------------------- tables


def _expand_grid(rows: list[list[tuple[str, int, int]]]) -> list[list[str]]:
    """Expand (text, rowspan, colspan) cells into a rectangular grid."""
    grid: list[list[str]] = []
    carry: dict[int, tuple[str, int]] = {}
    for row in rows:
        out: list[str] = []
        col = 0
        cells = list(row)
        while cells or any(c >= col for c in carry):
            if col in carry:
                text, left = carry[col]
                out.append(text)
                if left <= 1:
                    del carry[col]
                else:
                    carry[col] = (text, left - 1)
                col += 1
                continue
            if not cells:
                break
            text, rowspan, colspan = cells.pop(0)
            for offset in range(max(1, colspan)):
                out.append(text)
                if rowspan > 1:
                    carry[col + offset] = (text, rowspan - 1)
            col += max(1, colspan)
        grid.append(out)
    width = max((len(row) for row in grid), default=0)
    return [row + [""] * (width - len(row)) for row in grid]


def _looks_like_header(row: list[str], body: list[list[str]]) -> bool:
    filled = [cell for cell in row if cell]
    if len(filled) < 2 or not body:
        return False
    if any(len(cell) > 40 for cell in filled):
        return False
    numeric = sum(bool(re.fullmatch(r"[\d.:/\-\s%]+", cell)) for cell in filled)
    return numeric == 0


def table_blocks(grid: list[list[str]], table_id: str, *, header: list[str] | None = None, page: int | None = None, caption: str = "") -> list[RawBlock]:
    """Turn a rectangular grid into self-contained ``table_row`` blocks."""
    grid = [[normalize_text(cell) for cell in row] for row in grid]
    grid = [row for row in grid if any(row)]
    if not grid:
        return []
    width = max(len(row) for row in grid)
    if width == 1 or (len(grid) == 1 and header is None):
        # Layout tables and single cells are paragraphs, not data tables.
        return [RawBlock("paragraph", " ".join(c for c in row if c), page=page) for row in grid]
    body = grid
    if header is None and _looks_like_header(grid[0], grid[1:]):
        header, body = grid[0], grid[1:]
    header = [normalize_text(cell) for cell in (header or [])]
    blocks: list[RawBlock] = []
    for index, row in enumerate(body):
        # Merged cells repeat across columns; keep one copy in the row text.
        parts, last = [], None
        for col, cell in enumerate(row):
            if not cell or cell == last:
                last = cell or last
                continue
            last = cell
            name = header[col] if col < len(header) else ""
            parts.append(f"{name}: {cell}" if name and name != cell else cell)
        if not parts:
            continue
        text = " | ".join(parts)
        if caption:
            text = f"[{caption}] {text}"
        blocks.append(RawBlock("table_row", text, page=page, table={
            "table_id": table_id, "header": header, "row_index": index, "cells": row,
        }))
    return blocks


# ---------------------------------------------------------------- HTML


def _cell_text(cell: Tag) -> str:
    for br in cell.find_all("br"):
        br.replace_with(" ")
    return normalize_text(cell.get_text(" "))


def _html_table(table: Tag, table_id: str, caption: str = "") -> list[RawBlock]:
    raw_rows: list[list[tuple[str, int, int]]] = []
    header: list[str] | None = None
    for tr in table.find_all("tr"):
        if tr.find_parent("table") is not table:
            continue
        cells = [c for c in tr.find_all(["td", "th"], recursive=False)]
        row = []
        for cell in cells:
            try:
                rowspan = int(cell.get("rowspan", 1) or 1)
                colspan = int(cell.get("colspan", 1) or 1)
            except ValueError:
                rowspan = colspan = 1
            row.append((_cell_text(cell), min(rowspan, 50), min(colspan, 20)))
        if row:
            raw_rows.append(row)
    grid = _expand_grid(raw_rows)
    first = table.find("tr")
    if first is not None and (first.find_parent("thead") is not None or (first.find_all("th") and not first.find_all("td"))) and len(grid) > 1:
        header, grid = grid[0], grid[1:]
    return table_blocks(grid, table_id, header=header, caption=caption)


class _HtmlWalker:
    def __init__(self, table_prefix: str):
        self.blocks: list[RawBlock] = []
        self.buffer: list[str] = []
        self.table_prefix = table_prefix
        self.tables = 0

    def flush(self) -> None:
        text = normalize_text(" ".join(self.buffer))
        self.buffer = []
        if text:
            self.emit_text(text, "paragraph")

    def emit_text(self, text: str, kind: str, level: int = 0) -> None:
        text = normalize_text(text)
        if not text:
            return
        if kind == "paragraph":
            level = heading_level(text)
            if level:
                kind = "heading"
        self.blocks.append(RawBlock(kind, text, level=level))

    def walk(self, node: Tag) -> None:
        for child in node.children:
            if isinstance(child, NavigableString):
                if child.__class__.__name__ in ("Comment", "Doctype", "CData", "ProcessingInstruction"):
                    continue
                self.buffer.append(str(child))
                continue
            if not isinstance(child, Tag):
                continue
            name = child.name.lower()
            if name in SKIP_TAGS:
                continue
            if name not in BLOCK_TAGS:
                self.buffer.append(child.get_text(" "))
                continue
            self.flush()
            if name == "br" or name == "hr":
                continue
            if re.fullmatch(r"h[1-6]", name):
                self.emit_text(child.get_text(" "), "heading", int(name[1]))
            elif name == "table":
                if child.find("table") is not None:
                    self.walk_cells(child)
                else:
                    self.tables += 1
                    caption = child.find("caption")
                    caption_text = normalize_text(caption.get_text(" ")) if caption else ""
                    self.blocks.extend(_html_table(child, f"{self.table_prefix}t{self.tables}", caption_text))
            elif name == "li":
                nested = child.find_all(["ul", "ol", "table"], recursive=True)
                if nested:
                    clone = BeautifulSoup(str(child), "lxml").find("li")
                    for sub in clone.find_all(["ul", "ol", "table"]):
                        sub.decompose()
                    self.emit_text(clone.get_text(" "), "list_item")
                    for sub in child.find_all(["ul", "ol", "table"], recursive=False) or nested:
                        if sub.find_parent("li") is child:
                            self.walk(BeautifulSoup(f"<div>{sub}</div>", "lxml").div)
                else:
                    self.emit_text(child.get_text(" "), "list_item")
            elif name in ("p", "dt", "dd", "pre", "figcaption") and not child.find(["table", "div", "ul", "ol", "p"]):
                for br in child.find_all("br"):
                    br.replace_with("\n")
                for line in re.split(r"\n{1,}", child.get_text(" ") if name == "pre" else child.get_text("")):
                    self.emit_text(line, "paragraph")
            else:
                self.walk(child)
                self.flush()
        self.flush()

    def walk_cells(self, table: Tag) -> None:
        for cell in table.find_all(["td", "th"]):
            if cell.find("table") is None:
                self.walk(cell)


def html_blocks(root: Tag | str, table_prefix: str = "") -> list[RawBlock]:
    if isinstance(root, str):
        root = BeautifulSoup(root, "lxml").body or BeautifulSoup(root, "lxml")
    walker = _HtmlWalker(table_prefix)
    walker.walk(root)
    return [block for block in walker.blocks if block.text]


# ---------------------------------------------------------------- PDF

_OCR = None


class OcrDependencyMissing(RuntimeError):
    """The optional ``ocr`` extra (rapidocr-onnxruntime) is not installed."""


def ocr_image(image: bytes) -> str:
    """OCR a PNG byte string with RapidOCR (optional ``ocr`` dependency)."""
    global _OCR
    if _OCR is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError as exc:
            raise OcrDependencyMissing("rapidocr_onnxruntime not installed") from exc
        _OCR = RapidOCR()
    result, _ = _OCR(image)
    if not result:
        return ""
    return "\n".join(item[1] for item in result if item[2] >= 0.65)


def _line_blocks(text: str, page: int) -> Iterator[RawBlock]:
    """Merge wrapped PDF lines into paragraphs while keeping numbered items apart."""
    current: list[str] = []
    for line in text.splitlines():
        line = normalize_text(line)
        if not line:
            if current:
                yield RawBlock("paragraph", " ".join(current), page=page)
                current = []
            continue
        level = heading_level(line)
        starts_item = bool(re.match(r"^(?:第[一二三四五六七八九十百\d]+条|\d+[.、．)]|[（(]\d+[)）]|[•●▪\-]\s)", line))
        if level or starts_item:
            if current:
                yield RawBlock("paragraph", " ".join(current), page=page)
                current = []
            if level:
                yield RawBlock("heading", line, level=level, page=page)
                continue
        current.append(line)
        if line.endswith(("。", "；", ".", ";", "：", ":")) and len(" ".join(current)) > 80:
            yield RawBlock("paragraph", " ".join(current), page=page)
            current = []
    if current:
        yield RawBlock("paragraph", " ".join(current), page=page)


def pdf_extract(path: Path, *, ocr: bool = False, max_pages: int = 400, tables: bool = True) -> Extraction:
    import pymupdf

    result = Extraction(blocks=[])
    with pymupdf.open(path) as doc:
        result.pages = len(doc)
        if len(doc) > max_pages:
            result.warnings.append(f"truncated_pages:{len(doc)}")
        for number, page in enumerate(doc, 1):
            if number > max_pages:
                break
            table_rects = []
            if tables:
                try:
                    found = page.find_tables()
                    for index, table in enumerate(found.tables, 1):
                        grid = [[cell if cell is not None else "" for cell in row] for row in table.extract()]
                        if len(grid) < 2 or max(len(r) for r in grid) < 2:
                            continue
                        table_rects.append(pymupdf.Rect(table.bbox))
                        result.blocks.extend(table_blocks(grid, f"p{number}t{index}", page=number))
                except Exception as exc:  # pragma: no cover - parser robustness
                    result.warnings.append(f"table_error:p{number}:{type(exc).__name__}")
            if table_rects:
                pieces = []
                for x0, y0, x1, y1, text, *_ in page.get_text("blocks", sort=True):
                    rect = pymupdf.Rect(x0, y0, x1, y1)
                    if not any(rect.intersects(t) and (rect & t).get_area() > 0.5 * max(1, rect.get_area()) for t in table_rects):
                        pieces.append(text)
                text = "\n\n".join(pieces)
            else:
                text = page.get_text(sort=True)
            if len(normalize_text(text)) < 30 and not table_rects:
                if ocr:
                    result.ocr_attempted_pages.append(number)
                    try:
                        pix = page.get_pixmap(matrix=pymupdf.Matrix(1.8, 1.8), alpha=False)
                        text = ocr_image(pix.tobytes("png"))
                    except OcrDependencyMissing:
                        result.ocr_dependency_missing_pages.append(number)
                        result.warnings.append(f"ocr_dependency_missing:p{number}")
                        text = ""
                    except Exception as exc:  # pragma: no cover - OCR engine robustness
                        result.ocr_failed_pages.append(number)
                        result.warnings.append(f"ocr_failed:p{number}:{type(exc).__name__}")
                        text = ""
                    else:
                        if normalize_text(text):
                            result.ocr_succeeded_pages.append(number)
                            result.warnings.append(f"ocr:p{number}")
                        else:
                            result.warnings.append(f"ocr_empty:p{number}")
                else:
                    result.needs_ocr_pages.append(number)
            result.blocks.extend(_line_blocks(text, number))
    return result


# ---------------------------------------------------------------- Office


def docx_extract(path: Path) -> Extraction:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: list[RawBlock] = []
    tables = 0
    for element in doc.element.body:
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, doc)
            text = normalize_text(paragraph.text)
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "").casefold()
            match = re.search(r"(?:heading|标题)\s*(\d)", style)
            if match or style in ("title", "标题"):
                blocks.append(RawBlock("heading", text, level=int(match.group(1)) if match else 1))
            elif "list" in style:
                blocks.append(RawBlock("list_item", text))
            else:
                level = heading_level(text)
                blocks.append(RawBlock("heading" if level else "paragraph", text, level=level))
        elif element.tag.endswith("}tbl"):
            tables += 1
            table = Table(element, doc)
            grid = []
            for row in table.rows:
                try:
                    grid.append([cell.text for cell in row.cells])
                except Exception:  # malformed merged cells
                    continue
            blocks.extend(table_blocks(grid, f"t{tables}"))
    return Extraction(blocks=blocks, pages=1)


def xlsx_extract(path: Path, max_rows: int = 5000) -> Extraction:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[RawBlock] = []
    warnings: list[str] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets, 1):
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= max_rows:
                    warnings.append(f"truncated_rows:{sheet.title}")
                    break
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    rows.append(cells)
            if not rows:
                continue
            blocks.append(RawBlock("heading", normalize_text(sheet.title) or f"Sheet {sheet_index}", level=1))
            # Title rows above the real header ("ZJUI Class Schedule FA26") are
            # single-cell rows; keep them as paragraphs.
            while rows and sum(bool(c.strip()) for c in rows[0]) == 1 and len(rows) > 1:
                blocks.append(RawBlock("paragraph", normalize_text(next(c for c in rows[0] if c.strip()))))
                rows = rows[1:]
            width = max(len(r) for r in rows)
            grid = [r + [""] * (width - len(r)) for r in rows]
            keep = [i for i in range(width) if any(r[i].strip() for r in grid)]
            grid = [[r[i] for i in keep] for r in grid]
            if len(grid) > 1:
                blocks.extend(table_blocks(grid[1:], f"s{sheet_index}t1", header=grid[0]))
            else:
                blocks.extend(table_blocks(grid, f"s{sheet_index}t1"))
    finally:
        workbook.close()
    return Extraction(blocks=blocks, pages=len(workbook.sheetnames), warnings=warnings)


def text_extract(path: Path) -> Extraction:
    raw = path.read_bytes()
    for encoding in ("utf-8", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", "replace")
    return Extraction(blocks=list(_line_blocks(text, 1)), pages=1)


EXTRACTORS = {".pdf": pdf_extract, ".docx": docx_extract, ".xlsx": xlsx_extract, ".xlsm": xlsx_extract, ".txt": text_extract, ".md": text_extract}


def extract_file(path: Path, **options: Any) -> Extraction:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return pdf_extract(path, ocr=options.get("ocr", False))
    if suffix in EXTRACTORS:
        return EXTRACTORS[suffix](path)
    raise ValueError(f"unsupported:{suffix}")


def iter_text(blocks: Iterable[RawBlock]) -> str:
    return "\n".join(block.text for block in blocks)
