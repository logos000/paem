from pathlib import Path

import pytest

from structure.extract import (
    _expand_grid,
    docx_extract,
    heading_level,
    html_blocks,
    pdf_extract,
    table_blocks,
    xlsx_extract,
)


# --------------------------------------------------------------------- heading_level


def test_heading_level_chinese_chapter():
    assert heading_level("第一章 总则") == 1


def test_heading_level_chinese_numbered_list():
    assert heading_level("一、总则") == 2


def test_heading_level_chinese_parenthesized():
    assert heading_level("（一）概述") == 3


def test_heading_level_english_letter_heading():
    assert heading_level("A. Curriculum") == 2


def test_heading_level_english_part_chapter():
    assert heading_level("Part I Introduction") == 2
    assert heading_level("Chapter 2 Rules") == 2


def test_heading_level_zero_for_plain_sentence():
    assert heading_level("Not a heading, just a plain sentence.") == 0


def test_heading_level_zero_for_sentence_ending_with_full_stop_markers():
    assert heading_level("这是一句话。") == 0
    assert heading_level("这是一句话；") == 0


def test_heading_level_zero_for_overly_long_text():
    assert heading_level("一、" + "x" * 80) == 0


# --------------------------------------------------------------------- _expand_grid


def test_expand_grid_handles_rowspan_and_colspan():
    rows = [
        [("A", 2, 1), ("B", 1, 2)],
        [("C", 1, 1)],
    ]
    grid = _expand_grid(rows)
    assert grid == [["A", "B", "B"], ["A", "C", ""]]


def test_expand_grid_simple_rectangular_table():
    rows = [[("H1", 1, 1), ("H2", 1, 1)], [("a", 1, 1), ("b", 1, 1)]]
    assert _expand_grid(rows) == [["H1", "H2"], ["a", "b"]]


def test_expand_grid_empty_input():
    assert _expand_grid([]) == []


# --------------------------------------------------------------------- table_blocks


def test_table_blocks_width_one_becomes_paragraphs():
    grid = [["Just one column"], ["Another line"]]
    blocks = table_blocks(grid, "t1")
    assert [b.kind for b in blocks] == ["paragraph", "paragraph"]
    assert blocks[0].text == "Just one column"


def test_table_blocks_detects_header_from_row_text():
    grid = [["Name", "Time"], ["Alice", "8:00 pm"], ["Bob", "9:00 pm"]]
    blocks = table_blocks(grid, "t2")
    assert len(blocks) == 2
    assert blocks[0].text == "Name: Alice | Time: 8:00 pm"
    assert blocks[0].table == {"table_id": "t2", "header": ["Name", "Time"], "row_index": 0, "cells": ["Alice", "8:00 pm"]}
    assert blocks[1].text == "Name: Bob | Time: 9:00 pm"


def test_table_blocks_numeric_only_row_is_not_treated_as_header():
    grid = [["10", "20"], ["30", "40"]]
    blocks = table_blocks(grid, "t3")
    # no header detected: both rows become data rows (first with no header prefix)
    assert len(blocks) == 2
    assert blocks[0].table["header"] == []
    assert blocks[0].text == "10 | 20"


def test_table_blocks_merged_cell_dedupe_collapses_repeated_value():
    # "Seniors" spans the Class and Time columns after rowspan/colspan expansion
    grid = [["Class", "Time"], ["Seniors", "Seniors"]]
    blocks = table_blocks(grid, "t4")
    assert len(blocks) == 1
    assert blocks[0].text == "Class: Seniors"


def test_table_blocks_empty_grid_returns_nothing():
    assert table_blocks([], "t5") == []
    assert table_blocks([["", ""], ["", ""]], "t5") == []


def test_table_blocks_caption_is_prefixed():
    grid = [["Name", "Time"], ["Alice", "8pm"], ["Bob", "9pm"]]
    blocks = table_blocks(grid, "t6", caption="Schedule")
    assert blocks[0].text.startswith("[Schedule] ")


# --------------------------------------------------------------------- html_blocks


def test_html_blocks_heading_and_br_split_paragraph():
    html = "<div><h2>Overview</h2><p>Line one<br>Line two</p></div>"
    blocks = html_blocks(html)
    kinds_texts = [(b.kind, b.text) for b in blocks]
    assert ("heading", "Overview") in kinds_texts
    assert ("paragraph", "Line one") in kinds_texts
    assert ("paragraph", "Line two") in kinds_texts


def test_html_blocks_skips_script_tags():
    html = "<div><p>Visible text</p><script>var x = evil();</script></div>"
    blocks = html_blocks(html)
    texts = [b.text for b in blocks]
    assert texts == ["Visible text"]
    assert not any("evil" in t for t in texts)


def test_html_blocks_list_items_including_nested():
    html = """
    <ul>
      <li>Item A
        <ul><li>Nested A1</li><li>Nested A2</li></ul>
      </li>
      <li>Item B</li>
    </ul>
    """
    blocks = html_blocks(html)
    texts = [b.text for b in blocks if b.kind == "list_item"]
    assert texts == ["Item A", "Nested A1", "Nested A2", "Item B"]


def test_html_blocks_table_with_thead_extracts_header():
    html = """
    <table>
      <thead><tr><th>Name</th><th>Time</th></tr></thead>
      <tbody><tr><td>Alice</td><td>8pm</td></tr></tbody>
    </table>
    """
    blocks = html_blocks(html)
    rows = [b for b in blocks if b.kind == "table_row"]
    assert len(rows) == 1
    assert rows[0].text == "Name: Alice | Time: 8pm"
    assert rows[0].table["header"] == ["Name", "Time"]


def test_html_blocks_heading_levels_h1_to_h6():
    html = "<h1>Top</h1><h3>Sub</h3>"
    blocks = html_blocks(html)
    levels = {b.text: b.level for b in blocks}
    assert levels["Top"] == 1
    assert levels["Sub"] == 3


# --------------------------------------------------------------------- docx_extract


def test_docx_extract_headings_paragraphs_lists_and_tables(tmp_path):
    from docx import Document

    path = tmp_path / "doc.docx"
    doc = Document()
    doc.add_heading("Chapter One", level=1)
    doc.add_paragraph("This is a normal paragraph.")
    doc.add_paragraph("Bullet item", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Score"
    table.rows[1].cells[0].text = "Alice"
    table.rows[1].cells[1].text = "95"
    doc.save(str(path))

    extraction = docx_extract(path)
    kinds = [(b.kind, b.text) for b in extraction.blocks]
    assert ("heading", "Chapter One") in kinds
    assert ("paragraph", "This is a normal paragraph.") in kinds
    assert ("list_item", "Bullet item") in kinds
    table_rows = [b for b in extraction.blocks if b.kind == "table_row"]
    assert len(table_rows) == 1
    assert table_rows[0].text == "Name: Alice | Score: 95"


def test_docx_extract_heading_level_from_style_name(tmp_path):
    from docx import Document

    path = tmp_path / "doc2.docx"
    doc = Document()
    doc.add_heading("Level Two Heading", level=2)
    doc.save(str(path))

    extraction = docx_extract(path)
    heading = next(b for b in extraction.blocks if b.kind == "heading")
    assert heading.level == 2


# --------------------------------------------------------------------- xlsx_extract


def test_xlsx_extract_title_row_and_table(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append(["ZJUI Class Schedule FA26"])
    ws.append(["Course", "Time"])
    ws.append(["CS101", "8:00 pm"])
    ws.append(["CS102", "9:00 pm"])
    wb.save(str(path))

    extraction = xlsx_extract(path)
    kinds_texts = [(b.kind, b.text) for b in extraction.blocks]
    assert ("heading", "Schedule") in kinds_texts
    assert ("paragraph", "ZJUI Class Schedule FA26") in kinds_texts
    rows = [b for b in extraction.blocks if b.kind == "table_row"]
    assert rows[0].text == "Course: CS101 | Time: 8:00 pm"
    assert rows[1].text == "Course: CS102 | Time: 9:00 pm"
    assert extraction.pages == 1


def test_xlsx_extract_skips_empty_sheet(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "empty.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Empty"
    wb.create_sheet("AlsoEmpty")
    wb.save(str(path))

    extraction = xlsx_extract(path)
    assert extraction.blocks == []


# --------------------------------------------------------------------- pdf_extract


def test_pdf_extract_text_page_and_ocr_flag_for_blank_page(tmp_path):
    import pymupdf

    path = tmp_path / "doc.pdf"
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text(
        (72, 72),
        "This is the first page of the document with enough text content to pass the "
        "thirty character threshold easily.",
    )
    doc.new_page()  # second page has no text -> needs OCR
    doc.save(str(path))
    doc.close()

    extraction = pdf_extract(path, ocr=False)
    assert extraction.pages == 2
    assert extraction.needs_ocr_pages == [2]
    page_one_blocks = [b for b in extraction.blocks if b.page == 1]
    assert any("first page" in b.text for b in page_one_blocks)
    page_two_blocks = [b for b in extraction.blocks if b.page == 2]
    assert page_two_blocks == []


def test_pdf_extract_no_ocr_pages_when_all_pages_have_text(tmp_path):
    import pymupdf

    path = tmp_path / "doc2.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text(
        (72, 72),
        "A single page with plenty of real text content so it clears the OCR detection threshold.",
    )
    doc.save(str(path))
    doc.close()

    extraction = pdf_extract(path, ocr=False)
    assert extraction.needs_ocr_pages == []


def _two_page_pdf(tmp_path):  # page 1 has text, page 2 is blank
    import pymupdf

    path = tmp_path / "mixed.pdf"
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text(
        (72, 72),
        "Page one carries real text content that clears the OCR detection threshold.",
    )
    doc.new_page()
    doc.save(str(path))
    doc.close()
    return path


def test_pdf_extract_ocr_dependency_missing_degrades_per_page(tmp_path, monkeypatch):
    import sys

    monkeypatch.setattr("structure.extract._OCR", None)
    monkeypatch.setitem(sys.modules, "rapidocr_onnxruntime", None)  # import -> ImportError
    from structure.extract import OcrDependencyMissing

    extraction = pdf_extract(_two_page_pdf(tmp_path), ocr=True)
    # The missing dependency must not fail the whole document: page 1 text
    # survives, page 2 is accounted as dependency-missing.
    assert any("Page one" in b.text for b in extraction.blocks)
    assert extraction.ocr_attempted_pages == [2]
    assert extraction.ocr_dependency_missing_pages == [2]
    assert extraction.ocr_failed_pages == []
    assert "ocr_dependency_missing:p2" in extraction.warnings
    assert OcrDependencyMissing


def test_pdf_extract_ocr_engine_failure_degrades_per_page(tmp_path, monkeypatch):
    def boom(image):
        raise RuntimeError("engine crashed")

    monkeypatch.setattr("structure.extract.ocr_image", boom)
    extraction = pdf_extract(_two_page_pdf(tmp_path), ocr=True)
    assert any("Page one" in b.text for b in extraction.blocks)
    assert extraction.ocr_attempted_pages == [2]
    assert extraction.ocr_failed_pages == [2]
    assert extraction.ocr_dependency_missing_pages == []
    assert any(w.startswith("ocr_failed:p2:") for w in extraction.warnings)


def test_pdf_extract_ocr_empty_result_counts_attempted_only(tmp_path, monkeypatch):
    monkeypatch.setattr("structure.extract.ocr_image", lambda image: "")
    extraction = pdf_extract(_two_page_pdf(tmp_path), ocr=True)
    assert any("Page one" in b.text for b in extraction.blocks)
    assert extraction.ocr_attempted_pages == [2]
    assert extraction.ocr_succeeded_pages == []
    assert "ocr_empty:p2" in extraction.warnings
